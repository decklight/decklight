#!/usr/bin/env python3
# Copyright 2026 Gilles Philippart
# SPDX-License-Identifier: Apache-2.0
#
# Regenerates test/fixtures/shapes.pptx — the importer's shape fixture, written
# by python-pptx so it carries what PowerPoint itself writes: txBox="1" on text
# boxes, gallery presets, connectors snapped with stCxn/endCxn and one merely
# drawn, a group moved and resized after grouping (chOff/chExt), rotation, a
# flip, an unfilled region, a picture with arrows on it, freeform geometry, a
# table and speaker notes. Six slides, each a case test/import.test.mjs reads.
#
#   python3 -m venv .venv && .venv/bin/pip install python-pptx
#   .venv/bin/python test/fixtures/make-shapes.py test/fixtures/shapes.pptx
#
import struct, zlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

def png(path, w=320, h=180):
    # four flat bands: a picture that is recognisably one, in a few hundred bytes
    band = lambda y: [[40, 60, 160], [230, 90, 60], [60, 170, 110], [240, 200, 70]][y * 4 // h]
    rows = b''.join(b'\x00' + bytes(band(y) * w) for y in range(h))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    open(path, 'wb').write(b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b''))

def arrow(conn, end='tail', kind='triangle'):
    ln = conn.line._get_or_add_ln()
    e = etree.SubElement(ln, qn(f'a:{end}End')); e.set('type', kind)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
TITLE_ONLY, BLANK = prs.slide_layouts[5], prs.slide_layouts[6]

# 1 — a title slide with a text box: a layout, never a drawing
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Shapes, as PowerPoint writes them'
s.placeholders[1].text = 'a fixture for decklight import'
tb = s.shapes.add_textbox(Inches(1), Inches(5.5), Inches(6), Inches(1)); tb.text_frame.text = 'A text box (txBox="1"), beside a subtitle — this slide stays text.'

# 2 — a process: three chevrons, connectors SNAPPED to them, a table below
s = prs.slides.add_slide(TITLE_ONLY); s.shapes.title.text = 'How an order flows'
chev = []
for i, word in enumerate(['Plan', 'Build', 'Ship']):
    c = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(1 + i * 4), Inches(2), Inches(3), Inches(1.2)); c.text_frame.text = word; chev.append(c)
for a, b in [(0, 1), (1, 2)]:
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    conn.begin_connect(chev[a], 3); conn.end_connect(chev[b], 1); arrow(conn)
t = s.shapes.add_table(3, 3, Inches(1), Inches(4.2), Inches(8), Inches(1.5)).table
for r, row in enumerate([['step', 'owner', 'days'], ['Plan', 'PM', '3'], ['Ship', 'Ops', '1']]):
    for c, v in enumerate(row): t.cell(r, c).text = v
s.notes_slide.notes_text_frame.text = 'Three steps. The arrows are snapped, so this is a diagram even under --shapes strict.'

# 3 — an architecture in a GROUP that was then moved and resized; a loose line; rotation; an unfilled region
s = prs.slides.add_slide(TITLE_ONLY); s.shapes.title.text = 'Architecture'
g = s.shapes.add_group_shape()
api = g.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(2.5), Inches(1)); api.text_frame.text = 'API'
db = g.shapes.add_shape(MSO_SHAPE.CAN, Inches(5), Inches(2), Inches(1.5), Inches(1.4)); db.text_frame.text = 'Ledger'
dec = g.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(1.2), Inches(4), Inches(2), Inches(1.2)); dec.text_frame.text = 'valid?'
c1 = g.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0); c1.begin_connect(api, 3); c1.end_connect(db, 1); arrow(c1)
loose = g.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.2), Inches(3.05), Inches(2.2), Inches(3.95)); arrow(loose)   # drawn, never snapped
# the group is then dragged right and scaled up — its children keep their old numbers
g.left, g.top = Inches(4.5), Inches(1.6)
g.width, g.height = int(g.width * 1.3), int(g.height * 1.3)
region = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.4), Inches(3.5), Inches(4.8)); region.fill.background(); region.text_frame.text = 'one process'
tilt = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.4), Inches(2.4), Inches(1)); tilt.text_frame.text = 'worker'; tilt.rotation = 12
mirrored = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1.2), Inches(4.2), Inches(2), Inches(0.9)); mirrored.text_frame.text = 'back'
mirrored._element.spPr.xfrm.set('flipH', '1')

# 4 — an annotated screenshot: a picture, two loose arrows onto it, a callout
s = prs.slides.add_slide(TITLE_ONLY); s.shapes.title.text = 'The screen, annotated'
import tempfile, os
shot = os.path.join(tempfile.gettempdir(), 'decklight-fixture-shot.png')
png(shot)
pic = s.shapes.add_picture(shot, Inches(1), Inches(1.8), Inches(6.4), Inches(3.6))
for y in (Inches(2.6), Inches(4.4)):
    a = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.5), y, Inches(7.6), y); arrow(a)
call = s.shapes.add_shape(MSO_SHAPE.RECTANGULAR_CALLOUT, Inches(9.6), Inches(2.1), Inches(3), Inches(1)); call.text_frame.text = 'the button that matters'
lbl = s.shapes.add_textbox(Inches(9.6), Inches(4.1), Inches(3), Inches(0.8)); lbl.text_frame.text = 'and the total, here'

# 5 — hand-drawn geometry beside gallery presets, and a box with a real list in it
s = prs.slides.add_slide(TITLE_ONLY); s.shapes.title.text = 'Drawn by hand'
fb = s.shapes.build_freeform(Inches(1), Inches(3), scale=1.0)
fb.add_line_segments([(Inches(2.5), Inches(2)), (Inches(4), Inches(3)), (Inches(3.4), Inches(4.6)), (Inches(1.6), Inches(4.6))], close=True)
blob = fb.convert_to_shape(); blob.text_frame.text = 'blob'
star = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(5), Inches(2.2), Inches(1.6), Inches(1.6)); star.text_frame.text = 'wow'
hexa = s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(7), Inches(2.2), Inches(2), Inches(1.6)); hexa.text_frame.text = 'hex'
both = s.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, Inches(5), Inches(4.4), Inches(4), Inches(1)); both.text_frame.text = 'sync'
wordy = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), Inches(2.2), Inches(3.2), Inches(3.2))
tf = wordy.text_frame; tf.text = 'The service'
for line in ['owns the ledger', 'settles nightly', 'answers in under a second', 'and never loses a write']:
    p = tf.add_paragraph(); p.text = line; p.level = 1

# 6 — two columns of bullets in text boxes: a layout, and it must stay one
s = prs.slides.add_slide(TITLE_ONLY); s.shapes.title.text = 'Two columns'
for x, items in [(Inches(1), ['left one', 'left two']), (Inches(7), ['right one', 'right two'])]:
    b = s.shapes.add_textbox(x, Inches(2), Inches(5), Inches(3)); tf = b.text_frame; tf.text = items[0]
    p = tf.add_paragraph(); p.text = items[1]

import sys
out = sys.argv[1] if len(sys.argv) > 1 else 'shapes.pptx'
prs.save(out); print('wrote', out)
